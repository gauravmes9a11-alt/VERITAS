"""Script to generate a comprehensive, publication-quality DOCX file explaining
the Workplace Compliance Advisory Desk project architecture, mechanics, and operational workflow.
"""
import os
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for table cells."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tc_pr.append(tc_mar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    """Set subtle border styling for a table."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tbl_pr.append(borders)

def add_callout(doc, text, title="KEY INVARIANT", border_color="2B6CB0", bg_color="EDF2F7"):
    """Adds a stylish callout box with a colored left bar."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=200, right=150)
    
    # Left border only
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tc_pr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"[{title}] ")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    run_b = p.add_run(text)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    # Empty space after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

def build_docx(output_path: Path):
    doc = Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Color palette
    NAVY = RGBColor(0x1A, 0x36, 0x5D)
    SLATE = RGBColor(0x2B, 0x6C, 0xB0)
    CHARCOAL = RGBColor(0x2D, 0x37, 0x48)
    GREY = RGBColor(0x71, 0x80, 0x96)
    
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = CHARCOAL
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Workplace Compliance Advisory Desk")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(26)
    r_title.font.color.rgb = NAVY
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(18)
    r_sub = p_sub.add_run("System Architecture, Multi-Agent Cognitive Mechanics & Operational Technical Guide")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = SLATE

    # Meta bar table
    meta_tbl = doc.add_table(rows=1, cols=4)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_background(meta_tbl.cell(0, 0), "EDF2F7")
    set_cell_background(meta_tbl.cell(0, 1), "EDF2F7")
    set_cell_background(meta_tbl.cell(0, 2), "EDF2F7")
    set_cell_background(meta_tbl.cell(0, 3), "EDF2F7")
    set_table_borders(meta_tbl, color="CBD5E0")
    
    meta_items = [
        ("PROJECT ID", "Mini Project 09"),
        ("DOMAIN", "Indian Labour Law & HR"),
        ("TECH STACK", "FastAPI, LangGraph, RQ, FAISS"),
        ("VERIFICATION", "Multi-Stage Deterministic")
    ]
    for idx, (lbl, val) in enumerate(meta_items):
        cell = meta_tbl.cell(0, idx)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r1 = p.add_run(lbl + "\n")
        r1.font.size = Pt(8.5)
        r1.font.bold = True
        r1.font.color.rgb = GREY
        r2 = p.add_run(val)
        r2.font.size = Pt(9.5)
        r2.font.bold = True
        r2.font.color.rgb = NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(18)
        run.font.color.rgb = NAVY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(14)
        run.font.color.rgb = SLATE
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(12)
        run.font.color.rgb = CHARCOAL
        return p

    # --- SECTION 1 ---
    add_h1("1. Executive Summary & Problem Formulation")
    doc.add_paragraph(
        "The Workplace Compliance Advisory Desk is an enterprise-grade multi-agent AI system designed "
        "to deliver precise, legally grounded advisories on Indian employment and workplace compliance "
        "(focusing on Gratuity, Maternity Benefits, and Prevention of Workplace Sexual Harassment - POSH). "
        "Unlike generic LLM wrappers that generate plausible but legally hazardous advice, this desk is built "
        "under a strict zero-hallucination paradigm: every single factual and statutory assertion must cite an in-force, "
        "retrieved legal provision, and any uncertainty, missing fact, or conflict is decisively escalated to legal counsel."
    )
    
    add_callout(
        doc,
        "The service NEVER issues binding legal opinions or subjective advice. It either returns an "
        "'answer-with-citation' verified by deterministic code, requests missing establishment facts, "
        "or escalates with full context to human legal counsel. An unverified answer is NEVER delivered.",
        title="CORE PHILOSOPHY & ETHICAL GUARDRAIL"
    )

    add_h2("The Indian Labour Law Landscape & The 4.1 'Hard Case'")
    doc.add_paragraph(
        "A central motivation for this system is navigating the historic transition between the legacy labour acts "
        "(e.g., Payment of Gratuity Act, 1972; Maternity Benefit Act, 1961) and the new labour codes, primarily "
        "the Code on Social Security, 2020 (CoSS 2020), which was notified with commencement on 21 November 2025."
    )
    doc.add_paragraph(
        "Standard LLMs and naive RAG architectures catastrophically fail on the classic benchmark case from the project brief:\n"
        "Question: 'An employee engaged after 21 November 2025 on a written fixed-term contract of 18 months has completed the contract. Is gratuity payable?'"
    )
    doc.add_paragraph(
        "Why Naive Systems Fail:\n"
        "1. The General Rule Trap: Section 53(1) of CoSS 2020 states gratuity is payable on termination of employment after "
        "rendering continuous service for not less than 5 years. A standard retrieval system retrieves this rule and incorrectly answers 'No, gratuity is not payable because 18 months is less than 5 years'.\n"
        "2. The Hidden Proviso: The first proviso to Section 53(1) explicitly carves out fixed-term employees, establishing that continuous service of 5 years is NOT necessary for employees employed on fixed-term contracts.\n"
        "3. The Complete Verifier Solution: The Workplace Compliance Desk implements a dedicated section-aware chunker, "
        "a proactive Proviso-retrieval tool, and a 5-stage Verifier (Agent 3) that refuses to approve any advisory citing a statutory provision without having retrieved and analyzed all attached provisos and explanations."
    )

    # --- SECTION 2 ---
    add_h1("2. High-Level Architecture & Deployment Topology")
    doc.add_paragraph(
        "The system adheres to clean microservice boundaries designed for fault isolation, independent scaling, "
        "and reproducible latency. All microservices communicate over HTTP with strict JSON data contracts, "
        "coordinated asynchronously via Redis and RQ (Redis Queue)."
    )

    # Architecture Overview Table
    arch_tbl = doc.add_table(rows=1, cols=4)
    arch_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(arch_tbl)
    headers = ["Service Name", "Runtime / Framework", "Port / Protocol", "Primary Responsibility"]
    for i, h in enumerate(headers):
        cell = arch_tbl.cell(0, i)
        set_cell_background(cell, "1A365D")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9.5)

    services_data = [
        ("api", "FastAPI / Uvicorn", "Port 8000 (HTTP)", "Ingress validation, PII/toxicity guardrails, rate limiting, async job dispatch, status polling & trace delivery"),
        ("worker", "RQ / LangGraph", "Internal Consumer", "Executes the LangGraph coordinator state machine (`prep` -> `research` -> `advise` -> `finalize`)"),
        ("researcher", "FastAPI / FAISS / BM25", "Port 8001 (HTTP)", "Agent 1: Autonomous ReAct retrieval loop. Operates corpus search, section lookup, and provision expansion"),
        ("advisor", "FastAPI / Verifier", "Port 8002 (HTTP)", "Agent 2: Synthesizes structured assertions strictly from retrieved passages; executes the multi-stage Verifier"),
        ("redis", "Redis 7 / 8 (WSL / Linux)", "Port 6379 (TCP)", "Shared state store: job queues, run results, distributed circuit breakers, rate-limiting windows, and LRU caches"),
        ("stub / llm", "FastAPI / OpenAI spec", "Port 8080 (HTTP)", "Pluggable LLM provider (offline determinism stub or OpenAI/Anthropic/Ollama production API)"),
        ("indexer", "Python CLI / FAISS", "Batch / One-shot", "Hierarchical legal corpus parser, proviso embedding generator, and FAISS vector index builder"),
        ("observability", "Prometheus / Jaeger / Grafana", "Ports 9090 / 16686 / 3000", "End-to-end distributed tracing (OpenTelemetry), p50/p95 latency metrics, token cost tracking, and system health")
    ]

    for s_name, s_run, s_port, s_resp in services_data:
        row = arch_tbl.add_row()
        for c_idx, text in enumerate([s_name, s_run, s_port, s_resp]):
            cell = row.cells[c_idx]
            if s_name in ("api", "worker", "researcher", "advisor"):
                set_cell_background(cell, "F7FAFC")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.size = Pt(9)
            if c_idx == 0:
                run.bold = True
                run.font.color.rgb = NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- SECTION 3 ---
    add_h1("3. End-to-End Workflow & Cognitive Mechanics")
    doc.add_paragraph(
        "When an HR user or client submits a compliance question, the request undergoes an orchestrated journey "
        "spanning synchronous ingress protection, asynchronous queueing, and a multi-agent state machine."
    )

    add_h2("Phase 1: Ingress Guardrails & Validation Gate")
    doc.add_paragraph(
        "1. Schema Validation (app/contracts.py): Uses strict Pydantic schemas. Cross-field constraints are enforced:\n"
        "   - end_date must be greater than or equal to start_date.\n"
        "   - contract_months must reconcile (within +/- 1 month) with the date span.\n"
        "   - fixed_term contracts require explicit tenure data.\n"
        "2. PII Detection: Free-text inputs are scanned for phone numbers, email addresses, and 12-digit Indian national identity numbers (Aadhaar). Requests with PII are quarantined (HTTP 422) immediately to ensure privacy compliance.\n"
        "3. Prompt Injection Defense (app/guardrails.py): Scans for prompt override triggers ('ignore previous instructions', 'DAN mode', system prompt exfiltration). Blocks at score >= 3.0.\n"
        "4. Toxicity Filter: Rejects abusive or threatening language (threshold >= 2.0).\n"
        "5. Sliding Window Rate Limiting: Redis-backed sliding window throttles abuse per client ID."
    )

    add_h2("Phase 2: Asynchronous Job Enqueueing & Fast 202 Response")
    doc.add_paragraph(
        "Because multi-agent legal verification requires multiple LLM reasoning and retrieval iterations (taking between 5 to 25 seconds), "
        "the API does not block the client HTTP connection. Instead:\n"
        "- The API generates a unique run_id (UUID4) with an initialized correlation ID.\n"
        "- The job is enqueued in the Redis 'runs' queue via RQ.\n"
        "- The API immediately returns HTTP 202 Accepted with the run ID and a status URL: /v1/advisories/{run_id}.\n"
        "- The client UI or API caller polls the endpoint until status becomes 'completed'."
    )

    add_h2("Phase 3: The Coordinator State Machine (LangGraph)")
    doc.add_paragraph(
        "The worker process picks up the job and initializes the LangGraph State Graph (app/coordinator.py). "
        "The state machine executes the following deterministic transitions:"
    )

    add_h3("Step 3.1: Prep Node (Zero-LLM Cost Optimization)")
    doc.add_paragraph(
        "Before calling any LLM, the prep node executes fast rule-based classification:\n"
        "- Domain Classification: Categorizes into gratuity, maternity_benefit, workplace_harassment, or other.\n"
        "- Escalation Checks: Automatically escalates without spending LLM tokens if the query involves:\n"
        "  * Out-of-scope legal topics ('other') -> Route: legal-counsel-review\n"
        "  * Termination disputes (wrongful termination / labor court) -> Route: legal-counsel-review\n"
        "  * Live workplace harassment complaints -> Route: legal-counsel-review\n"
        "  * Service straddling commencement (service began prior to 21 Nov 2025 and ended after) -> Route: legal-counsel-review (due to complex transitional savings provisions)\n"
        "- Fact Verification: If essential parameters (e.g., employee count for maternity or committee status for POSH) are missing, "
        "it immediately transitions to finalize with route 'request-more-establishment-detail'."
    )

    add_h3("Step 3.2: Research Node (Researcher Agent Loop)")
    doc.add_paragraph(
        "If the query requires factual research, the coordinator delegates via HTTP to the Researcher Agent (app/agents/researcher_service.py). "
        "The Researcher executes a bounded ReAct cycle (up to 8 steps, max 10 tool calls, 90-second overall deadline):\n"
        "- Input Sanitization: User text is enclosed in untrusted tags: <untrusted_user_text>.\n"
        "- Autonomous Action Selection: The LLM selects tools from a strictly validated toolbox:\n"
        "  1) search_corpus: Multi-strategy retrieval (semantic vector search or BM25 lexical).\n"
        "  2) get_clause: Fetch a specific legal section by exact chunk identifier.\n"
        "  3) get_provisions: Fetch all provisos and explanations associated with a section reference.\n"
        "  4) get_request_field: Read structured facts from the original user submission.\n"
        "- Invariant Contract: The Researcher NEVER composes an advisory. Its sole output is an immutable bundle of retrieved passages."
    )

    add_h3("Step 3.3: Advise Node & The 5-Stage Complete Verifier")
    doc.add_paragraph(
        "The coordinator passes the passage bundle to the Advisor Agent (app/agents/advisor_service.py). "
        "The Advisor composes assertions strictly grounded in the provided passages. Then, before the advisory can be accepted, "
        "the deterministic Verifier (app/verify.py) audits every assertion against 5 statutory checks:"
    )

    # Verifier Checks Table
    v_tbl = doc.add_table(rows=1, cols=3)
    v_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(v_tbl)
    v_headers = ["Check Number & Name", "Audit Criterion", "System Action on Failure"]
    for i, h in enumerate(v_headers):
        cell = v_tbl.cell(0, i)
        set_cell_background(cell, "2B6CB0")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9.5)

    v_data = [
        ("Check 1: Citation Presence", "Every individual assertion must cite at least one explicit chunk ID.", "Rejected immediately. Marked 'uncited'."),
        ("Check 2: Citation Resolution", "Every cited chunk ID must exist in the retrieved passages bundle.", "Verification fails. Flags hallucinated citation ID."),
        ("Check 3: Temporal Validity (In-Force)", "The cited law must be in force on the query's as_of_date. (e.g. 1972 Act vs 2020 Code).", "Fails. Flags 'repealed_cited' and injects an actionable hint to search the superseded code."),
        ("Check 4: Textual & Numeric Support", "Lexical token overlap support >= 0.45. Every number in the claim MUST appear in the cited statutory text.", "Fails. Prevents numerical hallucination (e.g., inventing 26 weeks maternity or 15 days wages)."),
        ("Check 5: Completeness (Provisos & Explanations)", "Every proviso or explanation attached to a cited sub-section MUST have been retrieved and evaluated.", "CRITICAL: Generates an executable hint (e.g., 'call get_provisions(\"coss_2020:s53\")') and triggers a second research round!")
    ]

    for c_name, c_crit, c_fail in v_data:
        row = v_tbl.add_row()
        for idx, text in enumerate([c_name, c_crit, c_fail]):
            cell = row.cells[idx]
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(8.5)
            if idx == 0:
                run.bold = True
                run.font.color.rgb = NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_h3("Step 3.4: Feedback Loop & Second Research Round")
    doc.add_paragraph(
        "If Check 5 detects that the Researcher retrieved Section 53(1) (the 5-year gratuity rule) but missed the attached "
        "fixed-term proviso, the Verifier outputs a structured gap with an executable hint:\n"
        "  hint: 'coss_2020:s53(1) has provisos not yet read; call get_provisions(\"coss_2020:s53\") and reconsider'.\n"
        "The LangGraph coordinator loops back (advise -> research), feeding this hint directly into the Researcher. "
        "The Researcher fetches the proviso, the Advisor reads the exemption, and the Verifier finally approves the advisory!"
    )

    add_h3("Step 3.5: Degrade Mode (Resilience Fallback)")
    doc.add_paragraph(
        "If an LLM provider is down, an agent times out, or the distributed circuit breaker opens, the coordinator "
        "gracefully transitions to the degrade node:\n"
        "- Invokes /degraded_search on the researcher service: uses pure BM25 lexical retrieval without any LLM.\n"
        "- Packages the unverified raw statutory text as 'provisions_considered'.\n"
        "- Forces route to 'legal-counsel-review' with degraded=True.\n"
        "- Invariant guarantee: the system never fails silently or guesses."
    )

    add_h3("Step 3.6: Finalize Gate & Egress Guardrail")
    doc.add_paragraph(
        "At the final exit gate, before saving to Redis, the text is scanned by an output filter (app/guardrails.py):\n"
        "- Blocks forbidden legal opinion phrases (e.g., 'in my legal opinion', 'I hereby rule', 'you must fire', 'guaranteed to win').\n"
        "- If any forbidden phrase is detected, the advisory is suppressed, and the case is converted to an escalation."
    )

    # --- SECTION 4 ---
    add_h1("4. Legal Corpus Architecture & Ingestion Pipeline")
    doc.add_paragraph(
        "Legal texts cannot be chunked using standard 500-word sliding windows without destroying legal meaning. "
        "The desk features a purpose-built legal document pipeline (app/corpus/):"
    )
    doc.add_paragraph(
        "1. Hierarchical Decomposition (chunker.py):\n"
        "   - Act -> Section -> Sub-section -> Main Rule / Proviso / Explanation.\n"
        "   - Every chunk receives a deterministic global ID: {doc_id}:s{section}:{sub}:{kind}{index}\n"
        "     Example: coss_2020:s53:1:proviso1.\n"
        "2. Context-Prefixed Proviso Embeddings:\n"
        "   - Provisos frequently use relative language: 'Provided that nothing in this sub-section shall apply to...'. "
        "If embedded alone, the vector search cannot match questions asking about 'gratuity eligibility'.\n"
        "   - Solution: The chunker prefixes every proviso embedding with the main rule it qualifies: "
        "'[Code on Social Security | s.53(1) | Qualifies the rule that says: Gratuity shall be payable to an employee... -- Provided that...]'.\n"
        "3. Bi-Directional Metadata Linking:\n"
        "   - Main rule chunks store has_provisos=True and proviso_ids=[...].\n"
        "   - When an agent inspects the main rule, the system explicitly alerts it that exceptions exist.\n"
        "4. Temporal Indexing & Repeal Tracking:\n"
        "   - Documents specify in_force_from and in_force_until dates, plus superseded_by links.\n"
        "   - A query with as_of_date='2024-01-01' resolves the 1972 Gratuity Act as 'in_force' and the 2020 Code as 'pending_commencement'.\n"
        "   - A query with as_of_date='2026-01-01' resolves the 1972 Act as 'repealed' and CoSS 2020 as 'in_force'."
    )

    # --- SECTION 5 ---
    add_h1("5. Reliability, Caching & Distributed Resilience")
    doc.add_paragraph(
        "To satisfy production Service Level Objectives (SLOs) and control operational costs, the architecture incorporates "
        "industrial resilience patterns:"
    )
    doc.add_paragraph(
        "- Distributed Circuit Breaker (app/resilience.py): Researcher and Advisor microservices share circuit breaker states "
        "stored in Redis. If consecutive LLM calls fail past BREAKER_FAILURE_THRESHOLD (default 3), the breaker trips to OPEN for 30s. "
        "All worker threads immediately fast-fail to degrade mode without waiting for external API timeouts.\n"
        "- Three-Tiered Caching Hierarchy:\n"
        "  1. In-Memory LRU Cache: Caches dense vector embeddings in process.\n"
        "  2. Redis Retrieval Cache: Caches FAISS and BM25 query results keyed by (query, strategy, domain, index_version).\n"
        "  3. Whole-Run Result Cache: Caches completed verified responses keyed by normalized request facts + prompt version + model ID (1-hour TTL). "
        "Repeated inquiries cost ZERO LLM tokens and return in under 20 milliseconds!\n"
        "- Stateless Horizontal Scaling: RQ workers can be scaled horizontally (e.g. docker compose up -d --scale worker=4) "
        "to handle bursty traffic while sharing the same Redis broker and FAISS read-only corpus volume."
    )

    # --- SECTION 6 ---
    add_h1("6. Comprehensive API Reference")
    doc.add_paragraph("The desk exposes a modern RESTful API documented via OpenAPI / Swagger at http://localhost:8000/docs.")

    api_tbl = doc.add_table(rows=1, cols=3)
    api_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(api_tbl)
    a_headers = ["HTTP Method & Path", "Status Codes", "Description & Payload"]
    for i, h in enumerate(a_headers):
        cell = api_tbl.cell(0, i)
        set_cell_background(cell, "1A365D")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9.5)

    api_data = [
        ("POST /v1/advisories", "202 Accepted\n400 Guardrail\n422 Quarantine\n429 Rate Limit", "Submits compliance question. Scans PII/injection, validates cross-field constraints, enqueues run. Returns {run_id, status_url}."),
        ("GET /v1/advisories/{id}", "200 OK\n404 Not Found", "Polls run progress. Returns full result when completed: route, route_basis, summary, assertions with citations, missing_fields, escalation."),
        ("GET /v1/advisories/{id}/trace", "200 OK\n404 Not Found", "Reconstructs complete execution audit trail: all LLM steps, tool calls, retrieved passages, prompt hashes, latency, and verifier checks."),
        ("POST /v1/feedback", "201 Created\n404 Not Found", "Records human reviewer feedback: {run_id, rating: 'satisfactory' | 'unsatisfactory', reason: string}."),
        ("GET /v1/feedback/summary", "200 OK", "Aggregates satisfaction metrics and categorized feedback reasons across runs."),
        ("GET /v1/quality", "200 OK", "Returns real-time rolling evaluation scores: groundedness, faithfulness, context relevance, and hallucination rate."),
        ("GET /v1/quarantine", "200 OK", "Returns quarantined requests rejected by PII, toxic content, or invalid contract schemas."),
        ("GET /v1/contracts/{name}", "200 OK", "Delivers raw JSON Schema definitions for ask_request, run_result, or feedback."),
        ("GET /health", "200 OK\n503 Unavailable", "Service health check. Reports index loaded, corpus status (synthetic vs official), and Redis status."),
        ("GET /metrics", "200 OK", "Prometheus metrics endpoint: latency histograms, token counters, breaker states, and queue depth.")
    ]

    for m_path, m_stat, m_desc in api_data:
        row = api_tbl.add_row()
        for idx, text in enumerate([m_path, m_stat, m_desc]):
            cell = row.cells[idx]
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(8.5)
            if idx == 0:
                run.bold = True
                run.font.color.rgb = NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- SECTION 7 ---
    add_h1("7. Local & Production Runbook")
    doc.add_paragraph("The system can be operated in two environments:")

    add_h2("Local Development (Windows / Linux without Docker)")
    doc.add_paragraph(
        "Executed using the automated Python service manager scripts/run_services.py:\n"
        "1. Prerequisites: Python 3.10+ and Redis (native on Linux or via WSL on Windows).\n"
        "2. Command: python scripts/run_services.py (or .\\run_local.ps1 on Windows PowerShell).\n"
        "3. Automations Handled by Runner:\n"
        "   - Automatically detects Redis status on 127.0.0.1:6379; launches Redis inside WSL bound to 0.0.0.0 if inactive.\n"
        "   - Checks vector index manifest at data/index/main/manifest.json; runs the indexer automatically if absent.\n"
        "   - Concurrently spawns and monitors: stub (8080), researcher (8001), advisor (8002), api (8000), and worker (9101).\n"
        "   - Verifies health probes on all 5 services and streams logs to desk/logs/."
    )

    add_h2("Docker Compose (Production & Full Observability Stack)")
    doc.add_paragraph(
        "Deploy the full microservice and telemetry cluster:\n"
        "  docker compose up --build\n"
        "  Scale workers: docker compose up -d --scale worker=3\n"
        "Dashboard Interfaces:\n"
        "  - Web UI: http://localhost:8000\n"
        "  - Grafana Dashboards: http://localhost:3000 (admin / change-me)\n"
        "  - Jaeger Distributed Traces: http://localhost:16686\n"
        "  - Prometheus Metrics: http://localhost:9090"
    )

    # Save document
    doc.save(str(output_path))
    print(f"[+] DOCX successfully generated at {output_path}")

if __name__ == "__main__":
    out = Path("docs/Workplace_Compliance_Desk_Documentation.docx").resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    build_docx(out)
